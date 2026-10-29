#!/usr/bin/env python3
"""Run the fixed NOGA 2025 Laya pilot and write private review artifacts."""
from __future__ import annotations

import argparse
import csv
import json
import math
import resource
import statistics
import subprocess
import time
from collections import Counter, defaultdict
from pathlib import Path

import duckdb
import torch
from openpyxl import Workbook
from openpyxl.styles import Font
from laya import Router
from laya.lang import analyse

LEVELS = ("section", "division", "group", "class", "type")
QUESTION = "Which economic activity is the main activity described by this company purpose?"
REVIEW_COLUMNS = ["LAYA_SECTION_OK (j/n)", "LAYA_CODE_OK (j/n)", "RIGHT_CODE", "COMMENT"]


def read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def wilson(ok, n):
    if not n:
        return None, None
    z = 1.959963984540054
    p = ok / n
    d = 1 + z * z / n
    mid = (p + z * z / (2 * n)) / d
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return mid - half, mid + half


def swap_usage():
    try:
        return subprocess.check_output(["/usr/sbin/sysctl", "-n", "vm.swapusage"], text=True).strip()
    except (OSError, subprocess.CalledProcessError):
        return "unavailable"


def criteria_for(nodes):
    # Stable neutral keys keep the activity codes out of the model's choice labels.
    codes = list(nodes)
    keys = [chr(65 + i) if i < 26 else f"O{i:02d}" for i in range(len(codes))]
    return dict(zip(keys, (nodes[c]["LABEL_EN"] for c in codes))), dict(zip(keys, codes))


def _ranked_paths(paths):
    """Keep the two most likely paths; ties use the NOGA code for repeatability."""
    return sorted(paths, key=lambda path: (-path["score"], path["code"]))[:2]


def section_rotations(predict, purpose, names, section_map):
    """Run every S2 order and fail the aggregate if any rotation fails."""
    rotations, errors = [], []
    for shift in range(len(names)):
        question = {f"r{shift}": {"type": "choice", "instructions": QUESTION, "criteria": names,
                                  "option_order": [(p + shift) % len(names) for p in range(len(names))]}}
        started = time.perf_counter()
        try:
            answer = predict(purpose, question, f"r{shift}")
            mapping = dict(zip(names, section_map.values()))
            probabilities = {mapping[key]: value for key, value in answer.get("probabilities", {}).items()}
            choice = mapping.get(answer.get("choice"))
            elapsed = time.perf_counter() - started
            rotations.append((probabilities, elapsed))
            if not choice:
                errors.append("invalid_child")
            raw = {"options": list(section_map.values()), "probabilities": probabilities, "choice": choice,
                   "answer_confidence": answer.get("answer_confidence"), "seconds": elapsed,
                   "error": None if choice else "invalid_child"}
        except Exception as exc:
            elapsed = time.perf_counter() - started
            error = f"{type(exc).__name__}: {exc}"
            errors.append(error)
            raw = {"options": list(section_map.values()), "probabilities": {}, "choice": None,
                   "answer_confidence": None, "seconds": elapsed, "error": error}
        yield shift, raw, rotations, errors


def top2_paths(uid, purpose, section_record, children, infer):
    """Explore at most two paths, branching on uncertain sections and divisions."""
    probabilities = section_record.get("probabilities") or {}
    sections = sorted(probabilities.items(), key=lambda item: (-item[1], item[0]))
    if len(sections) < 2:
        return []
    section_uncertain = sections[0][1] - sections[1][1] < 0.15
    greedy_section = section_record.get("choice")
    if section_uncertain:
        starts = sections[:2]
    else:
        starts = [next((entry for entry in sections if entry[0] == greedy_section), sections[0])]
    paths = [{"code": code, "score": prob, "lineage": [("section", code)]}
             for code, prob in starts]
    for path in paths:
        choices = children.get(path["code"], [])
        if not choices:
            path["error"] = True
            continue
        if len(choices) == 1:
            path["division_options"] = [(choices[0], 1.0)]
            continue
        rec = infer(uid, "top2", "division", purpose, path["lineage"], choices)
        if rec["error"] or not rec["choice"]:
            path["error"] = True
            continue
        divisions = sorted(rec["probabilities"].items(), key=lambda item: (-item[1], item[0]))
        fork = len(divisions) > 1 and divisions[0][1] - divisions[1][1] < 0.15
        path["division_options"] = divisions[:2] if fork else divisions[:1]
    expanded = []
    for path in paths:
        for code, prob in path.pop("division_options", []):
            expanded.append({"code": code, "score": path["score"] * prob,
                             "lineage": path["lineage"] + [("division", code)]})
    paths = _ranked_paths(expanded)
    for level in LEVELS[2:]:
        next_paths = []
        for path in paths:
            choices = children.get(path["code"], [])
            if not choices:
                continue
            if len(choices) == 1:
                rec = {"choice": choices[0], "answer_confidence": 1.0, "error": None}
            else:
                rec = infer(uid, "top2", level, purpose, path["lineage"], choices)
            if rec["error"] or not rec["choice"]:
                continue
            prob = (rec.get("probabilities") or {}).get(rec["choice"], 1.0)
            next_paths.append({"code": rec["choice"], "score": path["score"] * prob,
                               "lineage": path["lineage"] + [(level, rec["choice"])]})
        paths = _ranked_paths(next_paths)
    return [path["code"] for path in paths if path["lineage"][-1][0] == "type"]


def top2_trigger(section_record, division_record):
    margins = []
    for record in (section_record, division_record):
        values = sorted((record or {}).get("probabilities", {}).values(), reverse=True)
        if len(values) > 1:
            margins.append(values[0] - values[1])
    return any(margin < 0.15 for margin in margins)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--database", required=True, type=Path,
                    help="DuckDB containing core_sample; pilot tables are written back to this file")
    ap.add_argument("--catalogue", required=True, type=Path)
    ap.add_argument("--xlsx", required=True, type=Path)
    ap.add_argument("--revision", required=True, help="Pinned Hugging Face commit SHA")
    args = ap.parse_args()

    input_db = duckdb.connect(str(args.database), read_only=True)
    cursor = input_db.execute("SELECT * FROM core_sample")
    sample_columns = [column[0] for column in cursor.description]
    companies = [dict(zip(sample_columns, row)) for row in cursor.fetchall()]
    input_db.close()
    catalogue_rows = read_csv(args.catalogue)
    nodes = {r["CODE"]: r for r in catalogue_rows}
    children = defaultdict(list)
    for r in catalogue_rows:
        children[r["PARENT_CODE"]].append(r["CODE"])
    for values in children.values():
        values.sort()

    device = "mps" if torch.backends.mps.is_available() else "cpu"
    router = Router(device=device, revision=args.revision, max_loaded=1)
    all_sections = {c: nodes[c] for c in children[""]}
    section_options, section_map = criteria_for(all_sections)
    run_started = time.perf_counter()
    t_load = time.perf_counter()
    router.load("multilingual")
    load_seconds = time.perf_counter() - t_load
    t_warm = time.perf_counter()
    router.predict("A company provides agricultural services.",
                   {"pick": {"type": "choice", "instructions": QUESTION,
                             "criteria": section_options}},
                   model="multilingual", max_len=1024, head_max_len=512)
    warmup_seconds = time.perf_counter() - t_warm
    raw = []
    by_uid = {c["UID"]: {"company": c, "levels": {}, "arms": {}, "alternatives": {}} for c in companies}

    def infer(uid, arm, level, purpose, ancestors, candidates, order=None):
        nonlocal load_seconds, warmup_seconds
        opts, key_to_code = criteria_for({code: nodes[code] for code in candidates})
        ancestry = "; ".join(f"{name} {code} ({nodes[code]['LABEL_EN']})" for name, code in ancestors)
        instruction = QUESTION + (f" Chosen ancestors: {ancestry}." if ancestry else "")
        question = {"pick": {"type": "choice", "instructions": instruction, "criteria": opts}}
        if order is not None:
            question["pick"]["option_order"] = order
        t0 = time.perf_counter()
        try:
            result = router.predict(purpose, question, model="multilingual", max_len=1024,
                                    head_max_len=512)
            answer = result["answers"]["pick"]
        except Exception as exc:
            elapsed = time.perf_counter() - t0
            record = {"uid": uid, "arm": arm, "level": level, "options": candidates,
                      "probabilities": {}, "choice": None, "answer_confidence": None,
                      "seconds": elapsed, "error": f"{type(exc).__name__}: {exc}"}
            raw.append(record)
            return record
        elapsed = time.perf_counter() - t0
        code = key_to_code.get(answer["choice"])
        record = {"uid": uid, "arm": arm, "level": level, "options": candidates,
                  "probabilities": {key_to_code[k]: v for k, v in answer.get("probabilities", {}).items()},
                  "choice": code, "answer_confidence": answer.get("answer_confidence"),
                  "seconds": elapsed, "error": None if code in candidates else "invalid_child"}
        raw.append(record)
        return record

    # Three independent section arms; S2 presents each option in each slot once.
    for idx, company in enumerate(companies, 1):
        uid, purpose = company["UID"], (company.get("REGISTER_PURPOSE") or "").strip()
        if not purpose:
            by_uid[uid]["no_purpose"] = True
            continue
        for arm, names in (("S1_EN", section_options), ("S2_EN_ROT", section_options),
                           ("S3_DE", {k: nodes[v]["LABEL_DE"] for k, v in section_map.items()})):
            if arm == "S2_EN_ROT":
                rotations, rotation_errors = [], []
                def predict_rotation(_purpose, question, key):
                    return router.predict(_purpose, question, model="multilingual", max_len=1024,
                                          head_max_len=512)["answers"][key]
                for shift, rotation, rotations, rotation_errors in section_rotations(
                        predict_rotation, purpose, names, section_map):
                    raw.append({"uid": uid, "arm": arm, "level": f"section_rotation_{shift:02d}", **rotation})
                averaged = ({code: statistics.mean(p.get(code, 0.0) for p, _ in rotations)
                             for code in section_map.values()} if not rotation_errors else {})
                code = max(averaged, key=averaged.get) if averaged else None
                rec = {"uid": uid, "arm": arm, "level": "section", "options": list(section_map.values()),
                       "probabilities": averaged, "choice": code,
                       "answer_confidence": averaged.get(code) if code else None,
                       "seconds": sum(t for _, t in rotations),
                       "error": "; ".join(rotation_errors) if rotation_errors else None}
                raw.append(rec)
            else:
                mapping = dict(zip(names, section_map.values()))
                t0 = time.perf_counter()
                try:
                    answer = router.predict(purpose, {"pick": {"type": "choice", "instructions": QUESTION,
                        "criteria": names}}, model="multilingual", max_len=1024, head_max_len=512)["answers"]["pick"]
                except Exception as exc:
                    elapsed = time.perf_counter() - t0
                    rec = {"uid": uid, "arm": arm, "level": "section", "options": list(section_map.values()),
                           "probabilities": {}, "choice": None, "answer_confidence": None,
                           "seconds": elapsed, "error": f"{type(exc).__name__}: {exc}"}
                    raw.append(rec); by_uid[uid]["arms"][arm] = rec
                    continue
                elapsed = time.perf_counter() - t0
                code = mapping.get(answer["choice"])
                rec = {"uid": uid, "arm": arm, "level": "section", "options": list(section_map.values()),
                       "probabilities": {mapping[k]: v for k, v in answer.get("probabilities", {}).items()},
                       "choice": code, "answer_confidence": answer.get("answer_confidence"),
                       "seconds": elapsed, "error": None if code else "invalid_child"}
                raw.append(rec)
            by_uid[uid]["arms"][arm] = rec
        if idx % 10 == 0:
            print(f"sections: {idx}/{len(companies)}; swap={swap_usage()}; mps_allocated={torch.mps.current_allocated_memory() if device == 'mps' else 0}", flush=True)

    comparisons = {a: sum(1 for c in companies if c.get("NOGA2025_SECTION") and
                          by_uid[c["UID"]]["arms"].get(a, {}).get("choice") == c["NOGA2025_SECTION"])
                   for a in ("S1_EN", "S2_EN_ROT", "S3_DE")}
    chosen_arm = max(comparisons, key=lambda a: (comparisons[a], a == "S2_EN_ROT"))

    # Greedy ladder and bounded two-path alternative, branching only at section/division.
    for idx, company in enumerate(companies, 1):
        uid, row = company["UID"], by_uid[company["UID"]]
        supplied_language = (company.get("REGISTER_PURPOSE_LANGUAGE") or "").strip().lower()
        row["language"] = (supplied_language if supplied_language and supplied_language != "unknown"
                           else analyse(company.get("REGISTER_PURPOSE") or "").get("language") or "unknown")
        purpose = (company.get("REGISTER_PURPOSE") or "").strip()
        if not purpose:
            continue
        section = row["arms"].get(chosen_arm, {}).get("choice")
        if not section:
            continue
        path, ancestors = section, [("section", section)]
        for level in LEVELS[1:]:
            candidates = children.get(path, [])
            if len(candidates) == 1:
                code = candidates[0]
                rec = {"uid": uid, "arm": "greedy", "level": level, "options": candidates,
                       "probabilities": {code: 1.0}, "choice": code, "answer_confidence": 1.0,
                       "seconds": 0.0, "error": None, "single_child": True}
                raw.append(rec)
            else:
                rec = infer(uid, "greedy", level, purpose, ancestors, candidates)
            row["levels"][level] = rec
            if rec["error"]:
                break
            path = rec["choice"]
            ancestors.append((level, path))
        row["levels"]["section"] = dict(row["arms"][chosen_arm])
        # Margin-triggered alternate section path is separate from the greedy result.
        greedy_division = row["levels"].get("division")
        if top2_trigger(row["arms"][chosen_arm], greedy_division):
            row["alternatives"]["type"] = top2_paths(uid, purpose, row["arms"][chosen_arm],
                                                       children, infer)
        if idx % 10 == 0:
            print(f"ladder: {idx}/{len(companies)}; swap={swap_usage()}; mps_allocated={torch.mps.current_allocated_memory() if device == 'mps' else 0}", flush=True)

    elapsed_total = time.perf_counter() - run_started
    headers = ["UID", "COMPANY_NAME", "REGISTER_PURPOSE", "DETECTED_LANGUAGE"]
    for level in LEVELS:
        headers += [f"LAYA_{level.upper()}_CODE", f"LAYA_{level.upper()}_LABEL_EN", f"LAYA_{level.upper()}_CONFIDENCE"]
    headers += ["TOP2_ALTERNATIVE", "CORTEX_NOGA2025_CODE", "CORTEX_NOGA2025_SECTION", "NOGA2008_GROUP",
                "NOGA2008_GROUP_LABEL", "section_agrees_cortex", "type_agrees_cortex"] + REVIEW_COLUMNS
    review_rows = []
    for company in companies:
        row = by_uid[company["UID"]]
        vals = [company.get("UID"), company.get("COMPANY_NAME"), company.get("REGISTER_PURPOSE"), row.get("language", "unknown")]
        for level in LEVELS:
            rec = row["levels"].get(level)
            code = rec.get("choice") if rec else None
            label = nodes.get(code, {}).get("LABEL_EN") if code else None
            vals += [str(code) if code else None, label, rec.get("answer_confidence") if rec else None]
        alt = row["alternatives"].get("type") or []
        vals += [", ".join(map(str, alt)) or None, company.get("NOGA2025_CODE"), company.get("NOGA2025_SECTION"),
                 company.get("NOGA2008_GROUP"), company.get("NOGA2008_GROUP_LABEL"),
                 _agreement(row["arms"].get(chosen_arm), company.get("NOGA2025_SECTION")),
                 _agreement(row["levels"].get("type"), company.get("NOGA2025_CODE"))] + [None] * len(REVIEW_COLUMNS)
        review_rows.append(vals)
    summary_rows = [["metric", "value"]]
    summary_rows.append(["companies", len(companies)]); summary_rows.append(["no_purpose", sum(not (c.get("REGISTER_PURPOSE") or "").strip() for c in companies)])
    summary_rows.append(["chosen_section_arm", chosen_arm])
    for arm, count in comparisons.items(): summary_rows.append([f"section_agreement_{arm}", f"{count}/{len(companies)}"])
    summary_rows.append(["language_counts", json.dumps(dict(Counter(by_uid[c["UID"]].get("language", "unknown") for c in companies)))])
    for level in LEVELS:
        counts = Counter()
        for company in companies:
            pred = by_uid[company["UID"]]["levels"].get(level)
            expected = _ancestor_code(company.get("NOGA2025_CODE"), level, nodes)
            if pred and pred.get("answer_confidence") is not None:
                confidence = float(pred["answer_confidence"])
                counts["confidence_<0.25" if confidence < .25 else
                       "confidence_0.25-0.49" if confidence < .5 else
                       "confidence_0.50-0.74" if confidence < .75 else "confidence_>=0.75"] += 1
            if pred and not pred.get("error") and expected and str(pred.get("choice")) == str(expected):
                counts["agree"] += 1
        low, high = wilson(counts["agree"], len(companies))
        summary_rows.append([f"{level}_agreement_cortex_wilson95", json.dumps({"agree": counts["agree"], "denominator": len(companies), "low": low, "high": high})])
        failures = sum(not (by_uid[c["UID"]]["levels"].get(level) and
                            not by_uid[c["UID"]]["levels"][level].get("error")) for c in companies)
        summary_rows.append([f"{level}_failures_no_purpose_in_denominator", json.dumps({"failures": failures, "no_purpose": sum(not (c.get("REGISTER_PURPOSE") or "").strip() for c in companies), "denominator": len(companies)})])
        summary_rows.append([f"{level}_confidence_distribution", json.dumps({k: v for k, v in counts.items() if k.startswith("confidence_")})])
        level_times = [r["seconds"] for r in raw if r["arm"] == "greedy" and r["level"] == level and r["seconds"] > 0]
        summary_rows.append([f"{level}_greedy_call_timing_seconds", json.dumps({"calls": len(level_times),
            "mean": statistics.mean(level_times) if level_times else None,
            "median": statistics.median(level_times) if level_times else None})])
    top2_types = {uid: row["alternatives"].get("type", []) for uid, row in by_uid.items()}
    top2_matches = sum(any(str(code) == str(company.get("NOGA2025_CODE")) for code in top2_types[company["UID"]])
                       for company in companies)
    top2_low, top2_high = wilson(top2_matches, len(companies))
    summary_rows.append(["top2_type_agreement_cortex_wilson95", json.dumps({"agree": top2_matches,
        "denominator": len(companies), "low": top2_low, "high": top2_high,
        "rows_with_type_path": sum(bool(v) for v in top2_types.values())})])
    summary_rows.append(["top2_prediction_calls", sum(r["arm"] == "top2" and r["seconds"] > 0 for r in raw)])
    prediction_calls = [r for r in raw if r["seconds"] > 0 and
                        not (r["arm"] == "S2_EN_ROT" and r["level"] == "section")]
    summary_rows.append(["actual_prediction_calls", len(prediction_calls)])
    summary_rows.append(["actual_prediction_mean_seconds", statistics.mean(r["seconds"] for r in prediction_calls)])
    for arm in ("S1_EN", "S2_EN_ROT", "S3_DE"):
        samples = [r["seconds"] for r in raw if r["arm"] == arm and
                   (r["level"] == "section" if arm != "S2_EN_ROT" else r["level"].startswith("section_rotation_"))]
        summary_rows.append([f"{arm}_call_timing_seconds", json.dumps({"calls": len(samples),
            "mean": statistics.mean(samples) if samples else None,
            "median": statistics.median(samples) if samples else None})])
    summary_rows.append(["per_company_seconds", elapsed_total / len(companies)])
    summary_rows.append(["load_seconds", load_seconds]); summary_rows.append(["warmup_seconds", warmup_seconds])
    summary_rows.append(["peak_rss_bytes", resource.getrusage(resource.RUSAGE_SELF).ru_maxrss])
    summary_rows.append(["mps_allocated_bytes", torch.mps.current_allocated_memory() if device == "mps" else None])
    summary_rows.append(["mps_driver_bytes", torch.mps.driver_allocated_memory() if device == "mps" else None])
    summary_rows.append(["swap_usage", swap_usage()])
    summary_rows.append(["extrapolated_1_1m_seconds", elapsed_total / len(companies) * 1100000])
    summary_rows.append(["checkpoint_revision", args.revision]); summary_rows.append(["device", device]); summary_rows.append(["total_seconds", elapsed_total])
    con = duckdb.connect(str(args.database))
    con.execute("CREATE OR REPLACE TABLE laya_calls(uid VARCHAR, arm VARCHAR, level VARCHAR, options JSON, probabilities JSON, choice VARCHAR, answer_confidence DOUBLE, timing_seconds DOUBLE, error VARCHAR)")
    con.executemany("INSERT INTO laya_calls VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)", [
        (r["uid"], r["arm"], r["level"], json.dumps(r["options"]), json.dumps(r["probabilities"]),
         r["choice"], r["answer_confidence"], r["seconds"], r["error"]) for r in raw])
    result_types = ["DOUBLE" if "_CONFIDENCE" in header else "VARCHAR" for header in headers]
    con.execute("CREATE OR REPLACE TABLE laya_result(%s)" % ", ".join(
        '"%s" %s' % (header.replace('"', '""'), typ) for header, typ in zip(headers, result_types)))
    placeholders = ", ".join("?" for _ in headers)
    con.executemany("INSERT INTO laya_result VALUES (%s)" % placeholders, review_rows)
    con.execute("CREATE OR REPLACE TABLE laya_summary(ordinal INTEGER, metric VARCHAR, value VARCHAR)")
    con.executemany("INSERT INTO laya_summary VALUES (?, ?, ?)", [
        (i, str(row[0]), None if row[1] is None else str(row[1])) for i, row in enumerate(summary_rows)])
    con.execute("DELETE FROM sources WHERE table_name IN ('laya_calls', 'laya_result', 'laya_summary')")
    con.executemany("INSERT INTO sources VALUES (?, ?, ?)", [
        ("laya_calls", "Laya pilot raw choice options, probabilities, choices, confidence, and timing", "noga_laya.duckdb"),
        ("laya_result", "Laya pilot company-level review output", "noga_laya.duckdb"),
        ("laya_summary", "Laya pilot aggregate counts and runtime measurements", "noga_laya.duckdb")])

    wb = Workbook()
    ws = wb.active; ws.title = "review"; ws.append(headers)
    for result_row in con.execute("SELECT * FROM laya_result ORDER BY UID").fetchall():
        ws.append(list(result_row))
    for cell in ws[1]: cell.font = Font(bold=True)
    for column in ws.iter_cols():
        if column[0].value and ("CODE" in str(column[0].value) or column[0].value in ("UID", "NOGA2008_GROUP")):
            for cell in column[1:]: cell.number_format = "@"
    summary = wb.create_sheet("summary")
    summary.append(["metric", "value"])
    for _, metric, value in con.execute("SELECT ordinal, metric, value FROM laya_summary ORDER BY ordinal").fetchall():
        summary.append([metric, value])
    con.close()
    args.xlsx.parent.mkdir(parents=True, exist_ok=True); wb.save(args.xlsx)
    print(json.dumps({"n": len(companies), "elapsed_seconds": elapsed_total, "load_seconds": load_seconds,
        "warmup_seconds": warmup_seconds, "device": device, "chosen_arm": chosen_arm, "section_agreement": comparisons,
        "peak_rss_bytes": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss, "swap_usage": swap_usage(),
        "mps_current_bytes": torch.mps.current_allocated_memory() if device == "mps" else None,
        "mps_driver_bytes": torch.mps.driver_allocated_memory() if device == "mps" else None,
        "calls": len(prediction_calls),
        "seconds_per_call_mean": statistics.mean(r["seconds"] for r in prediction_calls) if prediction_calls else None,
        "seconds_per_company": elapsed_total / len(companies),
        "extrapolated_1_1m_seconds": elapsed_total / len(companies) * 1100000}, indent=2))


def _agreement(pred, expected):
    if not pred or not expected or pred.get("error"):
        return None
    return "yes" if str(pred.get("choice")) == str(expected) else "no"


def _ancestor_code(type_code, level, nodes):
    if not type_code or type_code not in nodes:
        return None
    code = type_code
    while nodes.get(code, {}).get("LEVEL") != level:
        code = nodes.get(code, {}).get("PARENT_CODE")
        if not code or code not in nodes:
            return None
    return code


if __name__ == "__main__":
    main()
